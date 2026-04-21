#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations

"""
附件2通用性检验（独立脚本）
使用：Count-Layout Beam Heuristic for 3D Packing
适用设定：
1. 每种箱装产品数量相同，统一取 100 件；
2. 忽略重量相关约束；
3. 全部按标准件处理；
4. 只选典型封闭车型；
5. 问题1.1：单车满载率最大化；
6. 问题1.2：单车型下最少车辆数（通过重复调用问题1.1单车装载器实现）。
"""

import argparse
import csv
import random
import re
from collections import Counter
from dataclasses import dataclass
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl import load_workbook


@dataclass(frozen=True)
class Orientation:
    orient_id: int
    size: Tuple[int, int, int]
    rotation: str


@dataclass(frozen=True)
class CargoType:
    type_id: str
    name: str
    category: str
    length: int
    width: int
    height: int
    weight: float
    quantity: int
    stackable: bool
    can_rotate: bool
    fragile: bool
    oriented: bool
    max_support_pressure: float = 1e18

    @property
    def volume(self) -> int:
        return self.length * self.width * self.height


@dataclass(frozen=True)
class CargoItem:
    item_id: str
    type_id: str
    category: str
    original_size: Tuple[int, int, int]
    weight: float
    volume: int
    stackable: bool
    fragile: bool
    oriented: bool
    allowed_orientations: Tuple[Orientation, ...]


@dataclass
class PlacedCargo:
    item_id: str
    type_id: str
    category: str
    x: int
    y: int
    z: int
    length: int
    width: int
    height: int
    weight: float
    orientation_id: int
    rotation: str
    support_by: Optional[str] = None
    truck_id: Optional[str] = None

    @property
    def volume(self) -> int:
        return self.length * self.width * self.height


@dataclass(frozen=True)
class Truck:
    name: str
    length: int
    width: int
    height: int
    max_weight: float
    cost: float
    top_clearance: int = 0

    @property
    def effective_height(self) -> int:
        return self.height - self.top_clearance

    @property
    def effective_volume(self) -> int:
        return self.length * self.width * self.effective_height


@dataclass
class Space:
    x: int
    y: int
    z: int
    length: int
    width: int
    height: int

    @property
    def volume(self) -> int:
        return self.length * self.width * self.height


@dataclass
class LayoutState:
    spaces: List[Space]
    placed: List[PlacedCargo]
    remaining: List[CargoItem]
    used_weight: float
    used_volume: int
    score: float = 0.0
    packed_ratio: float = 0.0


TYPICAL_CLOSED_TRUCK_KEYWORDS = ("单节柜", "超高柜", "箱货", "集装箱")
EXCLUDE_TRUCK_KEYWORDS = ("高栏", "低栏", "板", "栏板", "平板")


def clean_text(x) -> str:
    if x is None:
        return ""
    return str(x).strip().replace("（", "(").replace("）", ")")


def parse_dimension_to_mm(value) -> Optional[int]:
    s = clean_text(value)
    if not s:
        return None
    nums = re.findall(r"\d+(?:\.\d+)?", s)
    if not nums:
        return None
    vals = [float(v) for v in nums]
    mx = max(vals)
    if "米" in s:
        return int(round(mx * 1000))
    return int(round(mx))


def load_products_from_excel(excel_path: str, quantity_each: int = 100) -> Dict[str, CargoType]:
    wb = load_workbook(excel_path, data_only=True)
    ws = wb["箱装产品尺寸"]
    rows = list(ws.iter_rows(values_only=True))
    products: Dict[str, CargoType] = {}
    pid = 1
    for row in rows:
        cells = [clean_text(x) for x in row]
        if not any(cells):
            continue
        name = cells[0]
        if not name or name in {"产品", "品名"}:
            continue
        parsed = [parse_dimension_to_mm(x) for x in row]
        nums = [x for x in parsed if x is not None]
        if len(nums) >= 6:
            dims = nums[-3:]
        elif len(nums) >= 3:
            dims = nums[:3]
        else:
            continue
        l, w, h = dims
        if min(l, w, h) <= 0:
            continue
        type_id = f"P{pid}"
        products[type_id] = CargoType(
            type_id=type_id,
            name=name,
            category="standard",
            length=l,
            width=w,
            height=h,
            weight=0.0,
            quantity=quantity_each,
            stackable=True,
            can_rotate=True,
            fragile=False,
            oriented=False,
        )
        pid += 1
    if not products:
        raise ValueError("未读取到任何有效产品数据")
    return products


def is_typical_closed_truck(name: str) -> bool:
    n = clean_text(name)
    if not n:
        return False
    if any(k in n for k in EXCLUDE_TRUCK_KEYWORDS):
        return False
    return any(k in n for k in TYPICAL_CLOSED_TRUCK_KEYWORDS)


def load_trucks_from_excel(excel_path: str) -> Dict[str, Truck]:
    wb = load_workbook(excel_path, data_only=True)
    ws = wb["车型尺寸"]
    rows = list(ws.iter_rows(values_only=True))
    trucks: Dict[str, Truck] = {}
    for row in rows:
        cells = [clean_text(x) for x in row]
        if not any(cells):
            continue
        name = cells[0]
        if not is_typical_closed_truck(name):
            continue
        parsed = [parse_dimension_to_mm(x) for x in row]
        nums = [x for x in parsed if x is not None]
        if len(nums) < 4:
            continue
        length, width, height = nums[0], nums[1], nums[2]
        cost = float(nums[-1])
        if min(length, width, height) <= 0:
            continue
        trucks[name] = Truck(
            name=name,
            length=length,
            width=width,
            height=height,
            max_weight=1e18,
            cost=cost,
        )
    if not trucks:
        raise ValueError("未读取到任何典型封闭车型")
    return trucks


def generate_orientations(cargo: CargoType) -> Tuple[Orientation, ...]:
    l, w, h = cargo.length, cargo.width, cargo.height
    dims = [
        (l, w, h, "LWH"),
        (l, h, w, "LHW"),
        (w, l, h, "WLH"),
        (w, h, l, "WHL"),
        (h, l, w, "HLW"),
        (h, w, l, "HWL"),
    ]
    uniq = []
    seen = set()
    for a, b, c, name in dims:
        key = (a, b, c)
        if key not in seen:
            seen.add(key)
            uniq.append(Orientation(len(uniq), key, name))
    return tuple(uniq)


def truck_score(truck: Truck, used_volume: int) -> Tuple[float, float, float]:
    space_util = used_volume / truck.effective_volume if truck.effective_volume > 0 else 0.0
    return space_util, 0.0, space_util


def expand_items(counts: Dict[str, int], cargo_types: Dict[str, CargoType], type_order: List[str]) -> List[CargoItem]:
    items: List[CargoItem] = []
    for type_id in type_order:
        cargo = cargo_types[type_id]
        n = counts.get(type_id, 0)
        if n <= 0:
            continue
        oris = generate_orientations(cargo)
        for i in range(1, n + 1):
            items.append(
                CargoItem(
                    item_id=f"{type_id}_{i:03d}",
                    type_id=type_id,
                    category=cargo.category,
                    original_size=(cargo.length, cargo.width, cargo.height),
                    weight=0.0,
                    volume=cargo.volume,
                    stackable=True,
                    fragile=False,
                    oriented=False,
                    allowed_orientations=oris,
                )
            )
    return items


def count_score(truck: Truck, counts: Dict[str, int], cargo_types: Dict[str, CargoType], type_order: List[str]) -> float:
    total_volume = sum(cargo_types[k].volume * v for k, v in counts.items())
    space_util, _, fullness = truck_score(truck, total_volume)
    if total_volume == 0:
        return -1e18
    vol_share = []
    for k in type_order:
        v = counts.get(k, 0) * cargo_types[k].volume
        if v > 0:
            vol_share.append(v / total_volume)
    concentration = sum(p * p for p in vol_share) if vol_share else 1.0
    return 0.95 * fullness - 0.03 * concentration


def sort_spaces_dblf(spaces: Sequence[Space]) -> List[Space]:
    return sorted(spaces, key=lambda s: (s.z, s.y, s.x, s.volume))


def remove_contained_spaces(spaces: List[Space]) -> List[Space]:
    result: List[Space] = []
    for i, a in enumerate(spaces):
        contained = False
        for j, b in enumerate(spaces):
            if i == j:
                continue
            if (
                a.x >= b.x and a.y >= b.y and a.z >= b.z and
                a.x + a.length <= b.x + b.length and
                a.y + a.width <= b.y + b.width and
                a.z + a.height <= b.z + b.height and
                (a.x, a.y, a.z, a.length, a.width, a.height) !=
                (b.x, b.y, b.z, b.length, b.width, b.height)
            ):
                contained = True
                break
        if not contained and a.length > 0 and a.width > 0 and a.height > 0:
            result.append(a)
    return result


def try_merge_two_spaces(a: Space, b: Space) -> Optional[Space]:
    if a.y == b.y and a.z == b.z and a.width == b.width and a.height == b.height:
        if a.x + a.length == b.x:
            return Space(a.x, a.y, a.z, a.length + b.length, a.width, a.height)
        if b.x + b.length == a.x:
            return Space(b.x, b.y, b.z, a.length + b.length, a.width, a.height)
    if a.x == b.x and a.z == b.z and a.length == b.length and a.height == b.height:
        if a.y + a.width == b.y:
            return Space(a.x, a.y, a.z, a.length, a.width + b.width, a.height)
        if b.y + b.width == a.y:
            return Space(b.x, b.y, b.z, a.length, a.width + b.width, a.height)
    if a.x == b.x and a.y == b.y and a.length == b.length and a.width == b.width:
        if a.z + a.height == b.z:
            return Space(a.x, a.y, a.z, a.length, a.width, a.height + b.height)
        if b.z + b.height == a.z:
            return Space(b.x, b.y, b.z, a.length, a.width, a.height + b.height)
    return None


def merge_spaces(spaces: List[Space]) -> List[Space]:
    spaces = remove_contained_spaces(spaces)
    changed = True
    while changed:
        changed = False
        n = len(spaces)
        for i in range(n):
            merged = False
            for j in range(i + 1, n):
                m = try_merge_two_spaces(spaces[i], spaces[j])
                if m is not None:
                    new_spaces = [spaces[k] for k in range(n) if k not in (i, j)]
                    new_spaces.append(m)
                    spaces = remove_contained_spaces(new_spaces)
                    changed = True
                    merged = True
                    break
            if merged:
                break
    return spaces


def best_grid_count(truck: Truck, cargo: CargoType) -> int:
    best = 0
    for ori in generate_orientations(cargo):
        l, w, h = ori.size
        nx = truck.length // l
        ny = truck.width // w
        nz = truck.effective_height // h
        best = max(best, nx * ny * nz)
    return best


def compute_upper_bounds(truck: Truck, cargo_types: Dict[str, CargoType]) -> Dict[str, int]:
    bounds: Dict[str, int] = {}
    for type_id, cargo in cargo_types.items():
        by_volume = truck.effective_volume // cargo.volume
        by_grid = best_grid_count(truck, cargo)
        bounds[type_id] = max(0, min(cargo.quantity, by_volume, by_grid))
    return bounds


def generate_count_candidates(truck: Truck, upper_bounds: Dict[str, int], cargo_types: Dict[str, CargoType], type_order: List[str], trials: int = 700, keep: int = 40, seed: int = 2026) -> List[Dict[str, int]]:
    rng = random.Random(seed)
    seen = set()
    pool: List[Tuple[float, Dict[str, int]]] = []
    for trial in range(trials):
        counts = {t: 0 for t in type_order}
        used_volume = 0
        attempts = 0
        while attempts < 4000:
            feasible: List[Tuple[float, str]] = []
            for t in type_order:
                if counts[t] >= upper_bounds[t]:
                    continue
                c = cargo_types[t]
                if used_volume + c.volume > truck.effective_volume:
                    continue
                new_counts = counts.copy()
                new_counts[t] += 1
                score = count_score(truck, new_counts, cargo_types, type_order) + rng.uniform(-0.01, 0.01)
                feasible.append((score, t))
            if not feasible:
                break
            feasible.sort(reverse=True)
            rcl_size = 2 + (trial % 4)
            chosen = rng.choice(feasible[: min(rcl_size, len(feasible))])[1]
            counts[chosen] += 1
            used_volume += cargo_types[chosen].volume
            attempts += 1
        key = tuple(counts[t] for t in type_order)
        if key not in seen and sum(counts.values()) > 0:
            seen.add(key)
            pool.append((count_score(truck, counts, cargo_types, type_order), counts))
    pool.sort(key=lambda x: x[0], reverse=True)
    return [counts for _, counts in pool[:keep]]


def build_item_orders(items: List[CargoItem], policy: int) -> List[CargoItem]:
    def ori_height(item: CargoItem) -> int:
        return max(o.size[2] for o in item.allowed_orientations)
    def base_area(item: CargoItem) -> int:
        return max(o.size[0] * o.size[1] for o in item.allowed_orientations)
    if policy == 0:
        return sorted(items, key=lambda it: (-it.volume, it.type_id, it.item_id))
    if policy == 1:
        return sorted(items, key=lambda it: (-base_area(it), -ori_height(it), -it.volume, it.item_id))
    if policy == 2:
        return sorted(items, key=lambda it: (-ori_height(it), -it.volume, it.item_id))
    return sorted(items, key=lambda it: (-it.volume, -base_area(it), it.item_id))


def can_place_in_space(item: CargoItem, ori: Orientation, space: Space) -> bool:
    l, w, h = ori.size
    return l <= space.length and w <= space.width and h <= space.height


def placement_gain(truck: Truck, state: LayoutState, item: CargoItem, ori: Orientation, space: Space) -> float:
    l, w, h = ori.size
    fill_ratio = (l * w * h) / max(1, space.length * space.width * space.height)
    low_z_bonus = 1.0 / (1.0 + space.z)
    tight_bonus = 1.0 / (1.0 + (space.length - l) + (space.width - w) + (space.height - h))
    used_volume = state.used_volume + item.volume
    space_util, _, fullness = truck_score(truck, used_volume)
    return fullness + 0.10 * fill_ratio + 0.03 * low_z_bonus + 0.02 * tight_bonus + 0.02 * space_util


def split_space(space: Space, placed: PlacedCargo) -> List[Space]:
    new_spaces: List[Space] = []
    if space.length > placed.length:
        new_spaces.append(Space(space.x + placed.length, space.y, space.z, space.length - placed.length, space.width, space.height))
    if space.width > placed.width:
        new_spaces.append(Space(space.x, space.y + placed.width, space.z, placed.length, space.width - placed.width, space.height))
    if space.height > placed.height:
        new_spaces.append(Space(space.x, space.y, space.z + placed.height, placed.length, placed.width, space.height - placed.height))
    return new_spaces


def apply_placement(truck: Truck, state: LayoutState, item_idx: int, ori: Orientation, space_idx: int) -> LayoutState:
    item = state.remaining[item_idx]
    space = state.spaces[space_idx]
    l, w, h = ori.size
    placed = PlacedCargo(
        item_id=item.item_id,
        type_id=item.type_id,
        category=item.category,
        x=space.x, y=space.y, z=space.z,
        length=l, width=w, height=h,
        weight=0.0,
        orientation_id=ori.orient_id,
        rotation=ori.rotation,
        truck_id=truck.name,
    )
    new_placed = list(state.placed)
    new_placed.append(placed)
    new_spaces = [s for idx, s in enumerate(state.spaces) if idx != space_idx]
    new_spaces.extend(split_space(space, placed))
    new_spaces = merge_spaces(new_spaces)
    new_spaces = sort_spaces_dblf(new_spaces)
    new_remaining = state.remaining[:item_idx] + state.remaining[item_idx + 1 :]
    used_volume = state.used_volume + item.volume
    space_util, _, fullness = truck_score(truck, used_volume)
    packed_ratio = len(new_placed) / max(1, len(new_placed) + len(new_remaining))
    return LayoutState(new_spaces, new_placed, new_remaining, 0.0, used_volume, fullness + 0.02 * packed_ratio + 0.01 * space_util, packed_ratio)


def pack_with_beam_search(truck: Truck, items: List[CargoItem], beam_width: int = 8, candidate_items_each_step: int = 6, candidate_moves_each_state: int = 12, policy: int = 0) -> LayoutState:
    ordered = build_item_orders(items, policy)
    init = LayoutState([Space(0, 0, 0, truck.length, truck.width, truck.effective_height)], [], ordered, 0.0, 0, 0.0, 0.0)
    beam: List[LayoutState] = [init]
    best = init
    while beam:
        next_beam: List[LayoutState] = []
        expanded_any = False
        for state in beam:
            if state.score > best.score or state.used_volume > best.used_volume:
                best = state
            if not state.remaining or not state.spaces:
                next_beam.append(state)
                continue
            item_candidates = state.remaining[: min(candidate_items_each_step, len(state.remaining))]
            move_pool: List[Tuple[float, int, Orientation, int]] = []
            for item_idx, item in enumerate(item_candidates):
                for space_idx, space in enumerate(state.spaces[: min(18, len(state.spaces))]):
                    for ori in item.allowed_orientations:
                        if can_place_in_space(item, ori, space):
                            gain = placement_gain(truck, state, item, ori, space)
                            move_pool.append((gain, item_idx, ori, space_idx))
            move_pool.sort(key=lambda x: x[0], reverse=True)
            move_pool = move_pool[:candidate_moves_each_state]
            if not move_pool:
                next_beam.append(state)
                continue
            expanded_any = True
            for _, item_idx, ori, space_idx in move_pool:
                next_beam.append(apply_placement(truck, state, item_idx, ori, space_idx))
        if not expanded_any:
            break
        next_beam.sort(key=lambda st: (st.score, st.used_volume, len(st.placed)), reverse=True)
        beam = next_beam[:beam_width]
        if beam and (beam[0].score > best.score or beam[0].used_volume > best.used_volume):
            best = beam[0]
    return best


def evaluate_candidate(truck: Truck, counts: Dict[str, int], cargo_types: Dict[str, CargoType], type_order: List[str], beam_width: int):
    items = expand_items(counts, cargo_types, type_order)
    best_state = None
    best_meta = {}
    for policy in range(4):
        state = pack_with_beam_search(truck, items, beam_width=beam_width, candidate_items_each_step=6, candidate_moves_each_state=12, policy=policy)
        if best_state is None or state.used_volume > best_state.used_volume:
            best_state = state
            best_meta = {"policy_index": policy}
    actual_counts = Counter(p.type_id for p in best_state.placed)
    target_counts = {k: v for k, v in counts.items() if v > 0}
    space_util, _, fullness = truck_score(truck, best_state.used_volume)
    packed_ratio = sum(actual_counts.values()) / max(1, sum(target_counts.values()))
    meta = {
        "truck": truck.name,
        "space_utilization": round(space_util, 6),
        "fullness_score": round(fullness, 6),
        "loaded_volume_mm3": int(best_state.used_volume),
        "target_counts": target_counts,
        "actual_counts": dict(actual_counts),
        "placed_item_count": len(best_state.placed),
        "packed_ratio": round(packed_ratio, 6),
        **best_meta,
    }
    return best_state, meta


def solve_problem_11(truck: Truck, cargo_types: Dict[str, CargoType], type_order: List[str], topk: int = 16, beam_width: int = 8, trials: int = 700, seed: int = 2026):
    upper_bounds = compute_upper_bounds(truck, cargo_types)
    candidates = generate_count_candidates(truck, upper_bounds, cargo_types, type_order, trials=trials, keep=max(topk * 3, 30), seed=seed)
    best_state = None
    best_meta = {}
    for idx, counts in enumerate(candidates[:topk]):
        state, meta = evaluate_candidate(truck, counts, cargo_types, type_order, beam_width=beam_width)
        meta["candidate_index"] = idx
        if best_state is None or (state.used_volume, meta.get("packed_ratio", 0.0)) > (best_state.used_volume, best_meta.get("packed_ratio", 0.0)):
            best_state, best_meta = state, meta
    return best_state, best_meta, upper_bounds


def create_remaining_cargo_types(base_cargo_types: Dict[str, CargoType], remaining: Dict[str, int]) -> Dict[str, CargoType]:
    new_types = {}
    for tid, cargo in base_cargo_types.items():
        q = remaining.get(tid, 0)
        new_types[tid] = CargoType(tid, cargo.name, cargo.category, cargo.length, cargo.width, cargo.height, 0.0, q, True, True, False, False)
    return new_types


def solve_problem_12_single_truck_type(truck: Truck, cargo_types: Dict[str, CargoType], type_order: List[str], topk: int = 12, beam_width: int = 6, trials: int = 500, seed: int = 2026):
    remaining = {tid: cargo.quantity for tid, cargo in cargo_types.items()}
    vehicle_results = []
    vehicle_idx = 1
    while sum(remaining.values()) > 0:
        local_cargo_types = create_remaining_cargo_types(cargo_types, remaining)
        state, meta, _ = solve_problem_11(truck, local_cargo_types, type_order, topk=topk, beam_width=beam_width, trials=trials, seed=seed + vehicle_idx)
        loaded = Counter(p.type_id for p in state.placed)
        if sum(loaded.values()) == 0:
            raise RuntimeError(f"{truck.name} 在第 {vehicle_idx} 辆车上未能装入任何货物")
        for tid, cnt in loaded.items():
            remaining[tid] = max(0, remaining[tid] - cnt)
        vehicle_results.append({
            "vehicle_id": f"{truck.name}_{vehicle_idx:03d}",
            "used_volume_mm3": state.used_volume,
            "space_utilization": round(state.used_volume / truck.effective_volume, 6),
            "actual_counts": dict(loaded),
            "placed": state.placed,
        })
        vehicle_idx += 1
    vehicle_count = len(vehicle_results)
    avg_space_util = sum(v["space_utilization"] for v in vehicle_results) / max(1, vehicle_count)
    summary = {
        "truck_name": truck.name,
        "vehicle_count": vehicle_count,
        "avg_space_utilization": round(avg_space_util, 6),
        "total_cost_per_1000km": round(vehicle_count * truck.cost, 2),
    }
    return vehicle_results, summary


def ensure_dir(path: str) -> None:
    import os
    os.makedirs(path, exist_ok=True)


def write_csv(path: str, header: List[str], rows: List[List[object]]) -> None:
    import os
    ensure_dir(os.path.dirname(path))
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def write_problem11_results(outdir: str, results: List[Dict[str, object]], cargo_types: Dict[str, CargoType]) -> None:
    rows = []
    for r in results:
        actual_named = {cargo_types[k].name: v for k, v in r["actual_counts"].items()}
        rows.append([
            r["truck_name"], r["truck_length_mm"], r["truck_width_mm"], r["truck_height_mm"],
            r["truck_volume_mm3"], r["used_volume_mm3"], r["space_utilization"],
            r["placed_item_count"], str(actual_named),
        ])
    import os
    write_csv(os.path.join(outdir, "problem11_single_truck_fullness.csv"),
              ["truck_name", "truck_length_mm", "truck_width_mm", "truck_height_mm", "truck_volume_mm3", "used_volume_mm3", "space_utilization", "placed_item_count", "actual_counts_named"],
              rows)


def write_problem12_results(outdir: str, summaries: List[Dict[str, object]]) -> None:
    rows = [[s["truck_name"], s["vehicle_count"], s["avg_space_utilization"], s["total_cost_per_1000km"]] for s in summaries]
    import os
    write_csv(os.path.join(outdir, "problem12_min_vehicle_count.csv"),
              ["truck_name", "vehicle_count", "avg_space_utilization", "total_cost_per_1000km"], rows)


def write_problem12_vehicle_details(outdir: str, truck_name: str, vehicle_results: List[Dict[str, object]], cargo_types: Dict[str, CargoType]) -> None:
    import os
    truck_dir = os.path.join(outdir, "problem12_vehicle_details", truck_name)
    ensure_dir(truck_dir)
    for vr in vehicle_results:
        rows = []
        for p in vr["placed"]:
            rows.append([p.item_id, p.type_id, cargo_types[p.type_id].name, p.x, p.y, p.z, p.length, p.width, p.height, p.orientation_id, p.rotation, p.truck_id])
        write_csv(os.path.join(truck_dir, f"{vr['vehicle_id']}_items.csv"),
                  ["item_id", "type_id", "product_name", "x", "y", "z", "length", "width", "height", "orientation_id", "rotation", "truck_id"], rows)


def parse_args():
    parser = argparse.ArgumentParser(description="附件2通用性检验：Count-Layout Beam Heuristic")
    parser.add_argument("--excel", type=str, default="附件2：验证数据集.xlsx", help="Excel 文件路径")
    parser.add_argument("--outdir", type=str, default="附件2_通用性检验结果_两阶段束搜索", help="输出目录")
    parser.add_argument("--quantity-each", type=int, default=100, help="每类产品统一数量")
    parser.add_argument("--seed", type=int, default=2026, help="随机种子")
    return parser.parse_args()


def main():
    import os
    args = parse_args()
    base_dir = os.path.dirname(os.path.abspath(__file__))
    excel_path = args.excel if os.path.isabs(args.excel) else os.path.join(base_dir, args.excel)
    outdir = args.outdir if os.path.isabs(args.outdir) else os.path.join(base_dir, args.outdir)

    cargo_types = load_products_from_excel(excel_path, quantity_each=args.quantity_each)
    trucks = load_trucks_from_excel(excel_path)
    type_order = list(cargo_types.keys())

    print(f"读取产品种类数: {len(cargo_types)}")
    print(f"读取典型封闭车型数: {len(trucks)}")
    print(f"产品列表: {[cargo_types[k].name for k in type_order]}")
    print(f"车型列表: {list(trucks.keys())}")

    # 只做问题1.1
    problem11_results = []
    for truck_name, truck in trucks.items():
        print(f"\n开始求解问题1.1：{truck.name}")
        state, meta, _ = solve_problem_11(
            truck,
            cargo_types,
            type_order,
            topk=6,
            beam_width=4,
            trials=120,
            seed=args.seed
        )
        result = {
            "truck_name": truck.name,
            "truck_length_mm": truck.length,
            "truck_width_mm": truck.width,
            "truck_height_mm": truck.height,
            "truck_volume_mm3": truck.effective_volume,
            "used_volume_mm3": state.used_volume,
            "space_utilization": round(state.used_volume / truck.effective_volume, 6),
            "placed_item_count": len(state.placed),
            "actual_counts": dict(Counter(p.type_id for p in state.placed)),
        }
        problem11_results.append(result)
        print(result)

    problem11_results.sort(key=lambda x: x["space_utilization"], reverse=True)
    write_problem11_results(outdir, problem11_results, cargo_types)

    print("\n问题1.1结果（前3名车型）:")
    for r in problem11_results[:3]:
        print(r)

    print(f"\n输出目录: {outdir}")

if __name__ == "__main__":
    main()
