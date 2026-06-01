import openpyxl
import pandas as pd
import numpy as np

class DataProcessor:
    def __init__(self, file_path):
        """初始化数据处理器
        
        Args:
            file_path (str): Excel文件路径
        """
        self.file_path = file_path
        # 使用data_only=True参数，这样openpyxl会读取公式计算后的值，而不是公式本身
        self.wb = openpyxl.load_workbook(file_path, data_only=True)
        
        # 初始化数据存储字典
        self.monthly_summary_data = {}
        self.silicon_cost_data = {}
        self.sales_data = {}
        self.variable_cost_data = {}
        self.silicon_reduction_data = {}
        
    def read_all_data(self):
        """读取所有数据"""
        print("正在读取所有表格...")
        self.read_monthly_summary_data()
        self.read_silicon_cost_data()
        self.read_sales_data()
        self.read_variable_cost_data()
        self.read_silicon_reduction_data()
        print("数据读取完成")
        
    def read_monthly_summary_data(self):
        """读取月度汇总数据"""
        summary_sheet = self.wb['汇总']
        
        # 读取销售收入数据 (B2到I2)
        self.monthly_summary_data['销售收入'] = [summary_sheet.cell(row=2, column=i).value for i in range(2, 10)]
        print(f"已读取 销售收入 数据: {self.monthly_summary_data['销售收入']}")
        
        # 读取生产变动成本 (B3到I3)
        self.monthly_summary_data['生产变动成本'] = [summary_sheet.cell(row=3, column=i).value for i in range(2, 10)]
        print(f"已读取 生产变动成本 数据: {self.monthly_summary_data['生产变动成本']}")
        
        # 读取生产公共成本 (B4到I4)
        self.monthly_summary_data['生产公共成本'] = [summary_sheet.cell(row=4, column=i).value for i in range(2, 10)]
        print(f"已读取 生产公共成本 数据: {self.monthly_summary_data['生产公共成本']}")
        
        # 读取人工成本 (B5到I5)
        self.monthly_summary_data['人工成本'] = [summary_sheet.cell(row=5, column=i).value for i in range(2, 10)]
        print(f"已读取 人工成本 数据: {self.monthly_summary_data['人工成本']}")
        
        # 读取折旧 (B6到I6)
        self.monthly_summary_data['折旧'] = [summary_sheet.cell(row=6, column=i).value for i in range(2, 10)]
        print(f"已读取 折旧 数据: {self.monthly_summary_data['折旧']}")
        
        # 读取营业税费 (B7到I7)
        self.monthly_summary_data['营业税费'] = [summary_sheet.cell(row=7, column=i).value for i in range(2, 10)]
        print(f"已读取 营业税费 数据: {self.monthly_summary_data['营业税费']}")
        
        # 读取销售费用 (B8到I8)
        self.monthly_summary_data['销售费用'] = [summary_sheet.cell(row=8, column=i).value for i in range(2, 10)]
        print(f"已读取 销售费用 数据: {self.monthly_summary_data['销售费用']}")
        
        # 读取管理费用 (B9到I9)
        self.monthly_summary_data['管理费用'] = [summary_sheet.cell(row=9, column=i).value for i in range(2, 10)]
        print(f"已读取 管理费用 数据: {self.monthly_summary_data['管理费用']}")
        
        # 读取财务费用 (B10到I10)
        self.monthly_summary_data['财务费用'] = [summary_sheet.cell(row=10, column=i).value for i in range(2, 10)]
        print(f"已读取 财务费用 数据: {self.monthly_summary_data['财务费用']}")
        
        # 读取销售利润 (B12到I12)
        self.monthly_summary_data['销售利润'] = [summary_sheet.cell(row=12, column=i).value for i in range(2, 10)]
        print(f"已读取 销售利润 数据: {self.monthly_summary_data['销售利润']}")
        
        # 读取所得税费用 (B13到I13)
        self.monthly_summary_data['所得税费用'] = [summary_sheet.cell(row=13, column=i).value for i in range(2, 10)]
        print(f"已读取 所得税费用 数据: {self.monthly_summary_data['所得税费用']}")
        
        # 读取销售净利润 (B14到I14)
        self.monthly_summary_data['销售净利润'] = [summary_sheet.cell(row=14, column=i).value for i in range(2, 10)]
        print(f"已读取 销售净利润 数据: {self.monthly_summary_data['销售净利润']}")
        
    def read_silicon_cost_data(self):
        """读取硅片成本数据"""
        silicon_sheet = self.wb['硅料单耗计算']
        
        # 读取1-8月的硅片成本数据
        for month in range(1, 9):
            row_offset = (month - 1) * 8 + 3  # 计算行偏移量
            
            # 读取N1、N2、N3和P的成本单价
            n1_cost = silicon_sheet.cell(row=row_offset, column=4).value
            n2_cost = silicon_sheet.cell(row=row_offset+1, column=4).value
            n3_cost = silicon_sheet.cell(row=row_offset+2, column=4).value
            p_cost = silicon_sheet.cell(row=row_offset+3, column=4).value
            
            self.silicon_cost_data[month] = {
                'N1': n1_cost,
                'N2': n2_cost,
                'N3': n3_cost,
                'P': p_cost
            }
            print(f"已读取 {month} 月硅片成本数据: {self.silicon_cost_data[month]}")
            
    def read_sales_data(self):
        """读取销售数据"""
        sales_sheet = self.wb['销售收入']
        
        # 读取1-8月的销售数据
        for month in range(1, 9):
            # 计算行偏移量
            if month == 1:
                row_offset = 3
            elif month == 2:
                row_offset = 10
            elif month == 3:
                row_offset = 17
            elif month == 4:
                row_offset = 24
            elif month == 5:
                row_offset = 31
            elif month == 6:
                row_offset = 38
            elif month == 7:
                row_offset = 45
            else:  # month == 8
                row_offset = 52
            
            # 读取N1、N2、N3和P的销售单价
            n1_price = sales_sheet.cell(row=row_offset, column=4).value
            n2_price = sales_sheet.cell(row=row_offset+1, column=4).value
            n3_price = sales_sheet.cell(row=row_offset+2, column=4).value
            p_price = sales_sheet.cell(row=row_offset+3, column=4).value
            
            # 读取N1、N2、N3和P的需求产量
            n1_demand = sales_sheet.cell(row=row_offset, column=5).value
            n2_demand = sales_sheet.cell(row=row_offset+1, column=5).value
            n3_demand = sales_sheet.cell(row=row_offset+2, column=5).value
            p_demand = sales_sheet.cell(row=row_offset+3, column=5).value
            
            self.sales_data[month] = {
                'price': {
                    'N1': n1_price,
                    'N2': n2_price,
                    'N3': n3_price,
                    'P': p_price
                },
                'demand': {
                    'N1': n1_demand,
                    'N2': n2_demand,
                    'N3': n3_demand,
                    'P': p_demand
                }
            }
            print(f"已读取 {month} 月销售数据")
            
    def read_variable_cost_data(self):
        """读取变动成本数据"""
        variable_cost_sheet = self.wb['生产变动成本']
        
        # 读取1-8月的变动成本数据
        for month in range(1, 9):
            # 计算行偏移量
            if month == 1:
                row_offset = 26
            elif month == 2:
                row_offset = 58
            elif month == 3:
                row_offset = 90
            elif month == 4:
                row_offset = 122
            elif month == 5:
                row_offset = 154
            elif month == 6:
                row_offset = 186  # 修正为正确的行号
            elif month == 7:
                row_offset = 218
            else:  # month == 8
                row_offset = 250
            
            # 读取N1、N2、N3和P的单位变动成本
            n1_cost = variable_cost_sheet.cell(row=row_offset, column=7).value  # G列
            n2_cost = variable_cost_sheet.cell(row=row_offset, column=9).value  # I列
            n3_cost = variable_cost_sheet.cell(row=row_offset, column=11).value  # K列
            p_cost = variable_cost_sheet.cell(row=row_offset, column=13).value  # M列
            
            self.variable_cost_data[month] = {
                'N1': n1_cost,
                'N2': n2_cost,
                'N3': n3_cost,
                'P': p_cost
            }
            print(f"已读取 {month} 月变动成本数据")
            
    def read_silicon_reduction_data(self):
        """读取硅泥核减数据"""
        variable_cost_sheet = self.wb['生产变动成本']
        
        # 读取1-8月的硅泥核减数据
        for month in range(1, 9):
            # 计算行偏移量
            if month == 1:
                row_offset = 30
            elif month == 2:
                row_offset = 62
            elif month == 3:
                row_offset = 94
            elif month == 4:
                row_offset = 126
            elif month == 5:
                row_offset = 158
            elif month == 6:
                row_offset = 190
            elif month == 7:
                row_offset = 222
            else:  # month == 8
                row_offset = 254
            
            # 读取硅泥核减数据
            reduction = variable_cost_sheet.cell(row=row_offset, column=13).value
            
            self.silicon_reduction_data[month] = reduction
            print(f"已读取 {month} 月硅泥核减数据: {reduction}")
            
    def get_average_values(self):
        """计算各项指标的平均值
        
        Returns:
            dict: 包含各项指标平均值的字典
        """
        averages = {}
        for name, values in self.monthly_summary_data.items():
            # 过滤掉None值并计算平均值
            valid_values = [v for v in values if v is not None]
            if valid_values:
                averages[name] = sum(valid_values) / len(valid_values)
            else:
                averages[name] = 0
        return averages

    def get_monthly_summary_data(self):
        """获取月度汇总数据，使用平均值进行预测
        
        Returns:
            dict: 按月份组织的汇总数据，使用平均值进行预测
        """
        # 获取各项指标的平均值
        averages = self.get_average_values()
        
        result = {}
        for month in range(1, 9):  # 1-8月
            month_data = {}
            for param_name, avg_value in averages.items():
                month_data[param_name] = avg_value
            result[month] = month_data
        return result
        
    def get_average_silicon_cost(self):
        """计算各类型硅片的平均成本
        
        Returns:
            dict: 包含各类型硅片平均成本的字典
        """
        averages = {'N1': 0, 'N2': 0, 'N3': 0, 'P': 0}
        counts = {'N1': 0, 'N2': 0, 'N3': 0, 'P': 0}
        
        for month_data in self.silicon_cost_data.values():
            for silicon_type in averages.keys():
                if month_data[silicon_type] is not None:
                    averages[silicon_type] += month_data[silicon_type]
                    counts[silicon_type] += 1
        
        for silicon_type in averages.keys():
            if counts[silicon_type] > 0:
                averages[silicon_type] /= counts[silicon_type]
        
        return averages

    def get_average_sales_data(self):
        """计算各类型硅片的平均销售数据
        
        Returns:
            dict: 包含各类型硅片平均价格和销量的字典
        """
        averages = {
            'price': {'N1': 0, 'N2': 0, 'N3': 0, 'P': 0},
            'demand': {'N1': 0, 'N2': 0, 'N3': 0, 'P': 0}
        }
        counts = {
            'price': {'N1': 0, 'N2': 0, 'N3': 0, 'P': 0},
            'demand': {'N1': 0, 'N2': 0, 'N3': 0, 'P': 0}
        }
        
        for month_data in self.sales_data.values():
            for silicon_type in ['N1', 'N2', 'N3', 'P']:
                if month_data['price'][silicon_type] is not None:
                    averages['price'][silicon_type] += month_data['price'][silicon_type]
                    counts['price'][silicon_type] += 1
                if month_data['demand'][silicon_type] is not None:
                    averages['demand'][silicon_type] += month_data['demand'][silicon_type]
                    counts['demand'][silicon_type] += 1
        
        for data_type in ['price', 'demand']:
            for silicon_type in averages[data_type].keys():
                if counts[data_type][silicon_type] > 0:
                    averages[data_type][silicon_type] /= counts[data_type][silicon_type]
        
        return averages

    def get_average_variable_cost(self):
        """计算各类型硅片的平均变动成本
        
        Returns:
            dict: 包含各类型硅片平均变动成本的字典
        """
        averages = {'N1': 0, 'N2': 0, 'N3': 0, 'P': 0}
        counts = {'N1': 0, 'N2': 0, 'N3': 0, 'P': 0}
        
        for month_data in self.variable_cost_data.values():
            for silicon_type in averages.keys():
                if month_data[silicon_type] is not None:
                    averages[silicon_type] += month_data[silicon_type]
                    counts[silicon_type] += 1
        
        for silicon_type in averages.keys():
            if counts[silicon_type] > 0:
                averages[silicon_type] /= counts[silicon_type]
        
        return averages

    def get_silicon_cost_data(self):
        """获取硅片成本数据（使用平均值）"""
        return {month: self.get_average_silicon_cost() for month in range(1, 9)}
        
    def get_sales_data(self):
        """获取销售数据（使用原始数据）"""
        return self.sales_data
        
    def get_variable_cost_data(self):
        """获取变动成本数据（使用原始数据）"""
        return self.variable_cost_data
        
    def get_silicon_reduction_data(self):
        """获取硅泥核减数据"""
        return self.silicon_reduction_data 
    
    def read_unit_variable_cost_data(self):
        """读取四种硅片单位变动成本数据"""
        print("正在读取四种硅片单位变动成本数据...")
        variable_cost_sheet = self.wb['生产变动成本']
        
        # 初始化存储单位变动成本的字典
        self.unit_variable_cost_data = {}
        
        # 定义每个月份对应的行号
        row_mapping = {
            1: 26,
            2: 58,
            3: 90,
            4: 122,
            5: 154,
            6: 186,
            7: 218,
            8: 250
        }
        
        # 定义列号映射
        column_mapping = {
            'N1': 7,  # G列
            'N2': 9,  # I列
            'N3': 11, # K列
            'P': 13   # M列
        }
        
        # 读取每个月份的四种硅片单位变动成本
        for month, row in row_mapping.items():
            month_data = {}
            for silicon_type, column in column_mapping.items():
                cost = variable_cost_sheet.cell(row=row, column=column).value
                month_data[silicon_type] = cost
            
            self.unit_variable_cost_data[month] = month_data
            print(f"已读取 {month} 月四种硅片单位变动成本数据: {month_data}")
        
        print("四种硅片单位变动成本数据读取完成")
    
    def get_unit_variable_cost_data(self):
        """获取四种硅片单位变动成本数据"""
        return self.unit_variable_cost_data